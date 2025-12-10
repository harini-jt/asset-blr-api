"""
Export Service Module
Handles CSV and Excel export functionality for various entities
"""

import csv
import io
from typing import List, Dict, Any, Optional
from datetime import datetime
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class ExportService:
    """Service for exporting data to CSV/Excel formats"""
    
    @staticmethod
    def generate_excel_response(
        data: List[Dict[str, Any]], 
        filename: str,
        fieldnames: Optional[List[str]] = None,
        sheet_name: str = "Data"
    ) -> StreamingResponse:
        """
        Generate an Excel file response from a list of dictionaries
        
        Args:
            data: List of dictionaries containing the data to export
            filename: Name of the file (without extension)
            fieldnames: Optional list of field names. If None, uses keys from first item
            sheet_name: Name of the worksheet
            
        Returns:
            StreamingResponse with Excel content
        """
        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        
        if not data:
            # Return empty Excel with headers only
            if fieldnames:
                ws.append(fieldnames)
                # Style header row
                for col_num, _ in enumerate(fieldnames, 1):
                    cell = ws.cell(row=1, column=col_num)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Save to bytes
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            return StreamingResponse(
                iter([output.getvalue()]),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={
                    "Content-Disposition": f"attachment; filename={filename}.xlsx"
                }
            )
        
        # Use provided fieldnames or extract from first item
        if fieldnames is None:
            fieldnames = list(data[0].keys())
        
        # Write header row
        ws.append(fieldnames)
        
        # Style header row
        for col_num, _ in enumerate(fieldnames, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Write data rows
        for row_data in data:
            row_values = [row_data.get(field, '') for field in fieldnames]
            ws.append(row_values)
        
        # Auto-adjust column widths
        for col_num, column_name in enumerate(fieldnames, 1):
            column_letter = get_column_letter(col_num)
            max_length = len(str(column_name))
            
            for row_num in range(2, ws.max_row + 1):
                cell_value = ws.cell(row=row_num, column=col_num).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            
            # Set width (with some padding)
            adjusted_width = min(max_length + 2, 50)  # Cap at 50
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Freeze header row
        ws.freeze_panes = 'A2'
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename={filename}.xlsx"
            }
        )
    
    @staticmethod
    def generate_csv_response(
        data: List[Dict[str, Any]], 
        filename: str,
        fieldnames: Optional[List[str]] = None
    ) -> StreamingResponse:
        """
        Generate a CSV file response from a list of dictionaries
        
        Args:
            data: List of dictionaries containing the data to export
            filename: Name of the file (without extension)
            fieldnames: Optional list of field names. If None, uses keys from first item
            
        Returns:
            StreamingResponse with CSV content
        """
        if not data:
            # Return empty CSV with headers only
            output = io.StringIO()
            if fieldnames:
                writer = csv.DictWriter(output, fieldnames=fieldnames)
                writer.writeheader()
            output.seek(0)
            return StreamingResponse(
                iter([output.getvalue()]),
                media_type="text/csv",
                headers={
                    "Content-Disposition": f"attachment; filename={filename}.csv"
                }
            )
        
        # Use provided fieldnames or extract from first item
        if fieldnames is None:
            fieldnames = list(data[0].keys())
        
        # Create CSV in memory
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(data)
        
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={
                "Content-Disposition": f"attachment; filename={filename}.csv"
            }
        )
    
    @staticmethod
    def format_work_order_for_export(work_order) -> Dict[str, Any]:
        """
        Format a WorkOrder object for CSV export
        
        Args:
            work_order: WorkOrder SQLModel instance
            
        Returns:
            Dictionary with formatted fields
        """
        return {
            'WO Number': work_order.wo_number,
            'Summary': work_order.summary,
            'Description': work_order.description or '',
            'Status': work_order.status,
            'Priority': work_order.priority,
            'Technician': work_order.technician or '',
            'Due Date': work_order.due_date.strftime('%Y-%m-%d') if work_order.due_date else '',
            'Time Spent (Hours)': work_order.time_spent_hours or 0,
            'Completion Notes': work_order.completion_notes or '',
            'Created At': work_order.created_at.strftime('%Y-%m-%d %H:%M:%S'),
        }
    
    @staticmethod
    def format_asset_for_export(asset) -> Dict[str, Any]:
        """
        Format an Asset object for CSV export
        
        Args:
            asset: Asset SQLModel instance
            
        Returns:
            Dictionary with formatted fields
        """
        return {
            'Asset ID': asset.asset_id,
            'Name': asset.name,
            'Category': asset.category,
            'Status': asset.status,
            'Serial Number': asset.serial_number or '',
            'Tag ID': asset.tag_id or '',
            'SAP ID': asset.sap_id or '',
            'Location ID': asset.location_id or '',
            'Station ID': asset.station_id or '',
            'Vendor Name': asset.vendor_name or '',
            'Purchase Date': asset.purchase_date.strftime('%Y-%m-%d') if asset.purchase_date else '',
            'Purchase Cost': asset.purchase_cost or 0,
            'Warranty Expiry': asset.warranty_expiry.strftime('%Y-%m-%d') if asset.warranty_expiry else '',
            'Invoice Number': asset.invoice_number or '',
            'Company Code': asset.company_code or '',
            'Plant Code': asset.plant_code or '',
            'Currency': asset.currency or '',
            'Meter Reading': asset.meter_reading or '',
            'Notes': asset.notes or '',
            'Physically Verified': 'Yes' if asset.physically_verified else 'No',
            'Created At': asset.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            'Updated At': asset.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
        }
    
    @staticmethod
    def format_inventory_for_export(item) -> Dict[str, Any]:
        """
        Format an Inventory item for CSV export
        
        Args:
            item: Inventory SQLModel instance
            
        Returns:
            Dictionary with formatted fields
        """
        return {
            'Item Name': item.item_name,
            'Part Number': item.part_number,
            'Description': item.description or '',
            'Stock on Hand': item.stock_on_hand,
            'Min Stock': item.min_stock,
            'Max Stock': item.max_stock,
            'Unit Cost': item.unit_cost or 0,
            'Book Value': item.book_value or 0,
            'Status': item.status,
            'Location ID': item.location_id or '',
            'Vendor ID': item.vendor_id or '',
            'Company Code': item.company_code or '',
            'Plant Code': item.plant_code or '',
            'Currency': item.currency or '',
            'Cost Center': item.cost_center or '',
            'Invoice Number': item.invoice_number or '',
            'Invoice Date': item.invoice_date.strftime('%Y-%m-%d') if item.invoice_date else '',
            'Capitalised On': item.capitalised_on.strftime('%Y-%m-%d') if item.capitalised_on else '',
            'State': item.state or '',
            'Remarks': item.remarks or '',
            'Physically Verified': 'Yes' if item.physically_verified else 'No',
            'Created At': item.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            'Updated At': item.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
        }
    
    @staticmethod
    def format_location_for_export(location) -> Dict[str, Any]:
        """
        Format a Location object for CSV export
        
        Args:
            location: Location SQLModel instance
            
        Returns:
            Dictionary with formatted fields
        """
        return {
            'ID': location.id,
            'Name': location.name,
            'Description': location.description or '',
            'Parent ID': location.parent_id or '',
        }
    
    @staticmethod
    def format_vendor_for_export(vendor) -> Dict[str, Any]:
        """
        Format a Vendor object for CSV export
        
        Args:
            vendor: Vendor SQLModel instance
            
        Returns:
            Dictionary with formatted fields
        """
        return {
            'ID': vendor.id,
            'Name': vendor.name,
            'Description': vendor.description or '',
            'Address': vendor.address or '',
            'Contact': vendor.contact or '',
            'Email': vendor.email or '',
            'Created At': vendor.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            'Updated At': vendor.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
        }


# Convenience functions for common export scenarios
def export_work_orders_csv(work_orders: List, filename: str = None) -> StreamingResponse:
    """
    Export work orders to CSV
    
    Args:
        work_orders: List of WorkOrder objects
        filename: Optional filename (default: work_orders_YYYY-MM-DD)
        
    Returns:
        StreamingResponse with CSV file
    """
    if filename is None:
        filename = f"work_orders_{datetime.now().strftime('%Y-%m-%d')}"
    
    formatted_data = [
        ExportService.format_work_order_for_export(wo) 
        for wo in work_orders
    ]
    
    return ExportService.generate_csv_response(formatted_data, filename)


def export_assets_csv(assets: List, filename: str = None) -> StreamingResponse:
    """
    Export assets to CSV
    
    Args:
        assets: List of Asset objects
        filename: Optional filename (default: assets_YYYY-MM-DD)
        
    Returns:
        StreamingResponse with CSV file
    """
    if filename is None:
        filename = f"assets_{datetime.now().strftime('%Y-%m-%d')}"
    
    formatted_data = [
        ExportService.format_asset_for_export(asset) 
        for asset in assets
    ]
    
    return ExportService.generate_csv_response(formatted_data, filename)


def export_inventory_csv(items: List, filename: str = None) -> StreamingResponse:
    """
    Export inventory items to CSV
    
    Args:
        items: List of Inventory objects
        filename: Optional filename (default: inventory_YYYY-MM-DD)
        
    Returns:
        StreamingResponse with CSV file
    """
    if filename is None:
        filename = f"inventory_{datetime.now().strftime('%Y-%m-%d')}"
    
    formatted_data = [
        ExportService.format_inventory_for_export(item) 
        for item in items
    ]
    
    return ExportService.generate_csv_response(formatted_data, filename)


def export_locations_csv(locations: List, filename: str = None) -> StreamingResponse:
    """
    Export locations to CSV
    
    Args:
        locations: List of Location objects
        filename: Optional filename (default: locations_YYYY-MM-DD)
        
    Returns:
        StreamingResponse with CSV file
    """
    if filename is None:
        filename = f"locations_{datetime.now().strftime('%Y-%m-%d')}"
    
    formatted_data = [
        ExportService.format_location_for_export(loc) 
        for loc in locations
    ]
    
    return ExportService.generate_csv_response(formatted_data, filename)


def export_vendors_csv(vendors: List, filename: str = None) -> StreamingResponse:
    """
    Export vendors to CSV
    
    Args:
        vendors: List of Vendor objects
        filename: Optional filename (default: vendors_YYYY-MM-DD)
        
    Returns:
        StreamingResponse with CSV file
    """
    if filename is None:
        filename = f"vendors_{datetime.now().strftime('%Y-%m-%d')}"
    
    formatted_data = [
        ExportService.format_vendor_for_export(vendor) 
        for vendor in vendors
    ]
    
    return ExportService.generate_csv_response(formatted_data, filename)


# Excel export functions
def export_work_orders_excel(work_orders: List, filename: str = None) -> StreamingResponse:
    """
    Export work orders to Excel
    
    Args:
        work_orders: List of WorkOrder objects
        filename: Optional filename (default: work_orders_YYYY-MM-DD)
        
    Returns:
        StreamingResponse with Excel file
    """
    if filename is None:
        filename = f"work_orders_{datetime.now().strftime('%Y-%m-%d')}"
    
    formatted_data = [
        ExportService.format_work_order_for_export(wo) 
        for wo in work_orders
    ]
    
    return ExportService.generate_excel_response(formatted_data, filename, sheet_name="Work Orders")


def export_assets_excel(assets: List, filename: str = None) -> StreamingResponse:
    """
    Export assets to Excel
    
    Args:
        assets: List of Asset objects
        filename: Optional filename (default: assets_YYYY-MM-DD)
        
    Returns:
        StreamingResponse with Excel file
    """
    if filename is None:
        filename = f"assets_{datetime.now().strftime('%Y-%m-%d')}"
    
    formatted_data = [
        ExportService.format_asset_for_export(asset) 
        for asset in assets
    ]
    
    return ExportService.generate_excel_response(formatted_data, filename, sheet_name="Assets")


def export_inventory_excel(items: List, filename: str = None) -> StreamingResponse:
    """
    Export inventory items to Excel
    
    Args:
        items: List of Inventory objects
        filename: Optional filename (default: inventory_YYYY-MM-DD)
        
    Returns:
        StreamingResponse with Excel file
    """
    if filename is None:
        filename = f"inventory_{datetime.now().strftime('%Y-%m-%d')}"
    
    formatted_data = [
        ExportService.format_inventory_for_export(item) 
        for item in items
    ]
    
    return ExportService.generate_excel_response(formatted_data, filename, sheet_name="Inventory")


def export_locations_excel(locations: List, filename: str = None) -> StreamingResponse:
    """
    Export locations to Excel
    
    Args:
        locations: List of Location objects
        filename: Optional filename (default: locations_YYYY-MM-DD)
        
    Returns:
        StreamingResponse with Excel file
    """
    if filename is None:
        filename = f"locations_{datetime.now().strftime('%Y-%m-%d')}"
    
    formatted_data = [
        ExportService.format_location_for_export(loc) 
        for loc in locations
    ]
    
    return ExportService.generate_excel_response(formatted_data, filename, sheet_name="Locations")


def export_vendors_excel(vendors: List, filename: str = None) -> StreamingResponse:
    """
    Export vendors to Excel
    
    Args:
        vendors: List of Vendor objects
        filename: Optional filename (default: vendors_YYYY-MM-DD)
        
    Returns:
        StreamingResponse with Excel file
    """
    if filename is None:
        filename = f"vendors_{datetime.now().strftime('%Y-%m-%d')}"
    
    formatted_data = [
        ExportService.format_vendor_for_export(vendor) 
        for vendor in vendors
    ]
    
    return ExportService.generate_excel_response(formatted_data, filename, sheet_name="Vendors")
